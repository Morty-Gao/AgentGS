import gzip
import hashlib
import importlib
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from agentgs.datasets.parsers.energy import process_gip
from agentgs.datasets.references import load_reference

ROOT = Path(__file__).parents[3]


def test_workbook_export_preserves_sheets_values_formulas_and_energy_results(monkeypatch, tmp_path):
    monkeypatch.syspath_prepend(str(ROOT / "scripts/prepare"))
    converter = importlib.import_module("_gem")
    workbook = Workbook()
    facilities = workbook.active
    facilities.title = "Power facilities"
    facilities.append(
        [
            "Type",
            "Country/area",
            "Plant / Project name",
            "Capacity (MW)",
            "Status",
            "Latitude",
            "Longitude",
        ]
    )
    facilities.append(["Solar", "China", "太阳能电站", 12.5, "operating", 20.0, 100.0])
    about = workbook.create_sheet("About")
    about.append(["source notes", True, None, 0])
    regions = workbook.create_sheet("Regions, area, and countries")
    regions.append(["country", "code", "formula"])
    regions.append(["China", "CHN", "=1+1"])
    original = tmp_path / "source.xlsx"
    workbook.save(original)
    workbook.close()
    target = converter.convert_workbook(original, tmp_path / "gem.jsonl")
    records = [json.loads(line) for line in target.read_text().splitlines()]
    assert records[0]["sheets"] == ["Power facilities", "About", "Regions, area, and countries"]
    assert records[0]["source_sha256"] == hashlib.sha256(original.read_bytes()).hexdigest()
    notes = next(r for r in records if r.get("sheet") == "About")
    assert notes["values"] == ["source notes", True, None, 0]
    formula = next(r for r in records if r.get("sheet") == regions.title and r.get("row") == 2)
    assert formula["formulas"]["C2"]["text"] == "1+1"
    assert formula["formulas"]["C2"]["cached_value"] is None
    packed = converter.compress_jsonl(target, tmp_path / "gem.jsonl.gz")
    with gzip.open(packed, "rb") as stream:
        assert stream.read() == target.read_bytes()
    expected = process_gip(original)
    for actual in [process_gip(target), process_gip(packed)]:
        for before, after in zip(expected, actual, strict=True):
            pd.testing.assert_frame_equal(before, after)


def test_reference_values_are_external_inputs():
    energy = load_reference("reference/energy.json")
    assert energy["capitals"]["DE"][:2] == [52.52, 13.405]
    assert len(energy["capacity_series"]) == 10
    assert len(load_reference("reference/ocean.json")["ocean_boxes"]) == 10
    assert load_reference("reference/geography.json")["mean_earth_radius_km"] == 6371.0088
    assert len(load_reference("geoquestions/profile.json")["source_ids"]) == 40
    assert not (ROOT / "src/agentgs/datasets/geoquestions_snapshot.json").exists()
