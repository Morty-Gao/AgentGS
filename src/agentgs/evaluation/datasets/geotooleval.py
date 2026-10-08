from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CATEGORY_NAMES = {
    1: "DeepEarthTools",
    2: "GeoBigDataTools",
    3: "MineralResourcesTools",
    4: "ClimateChangeTools",
    5: "EnergyResourcesTools",
    6: "OceanCarbonTools",
    7: "DeepTimeBiosphereTools",
    8: "CCSTools",
    9: "RemoteSensingTools",
    10: "NaturalHazardTools",
}


@dataclass(frozen=True)
class EvalTask:
    question_id: int
    category_id: int
    category: str
    question: str
    expected_answer: str
    answer_type: str
    input_resource: str
    chain_mode: str
    gold_chain_text: str
    gold_steps: tuple[dict[str, Any], ...]


def load_tasks(workbook_path: Path) -> list[EvalTask]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["评测集合"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        columns = {name: index for index, name in enumerate(headers)}
        required = {
            "题目序号",
            "所属类别",
            "题目正文",
            "题目答案",
            "标准工具链调用",
            "单/多工具链调用",
            "答案类型",
            "输入资源",
            "标准工具调用明细",
        }
        missing = required - columns.keys()
        if missing:
            raise ValueError(f"GeoToolEval workbook is missing columns: {sorted(missing)}")
        tasks: list[EvalTask] = []
        for row in rows:
            question_id = int(row[columns["题目序号"]])
            category_id = int(row[columns["所属类别"]])
            raw_steps = row[columns["标准工具调用明细"]]
            try:
                gold_steps = json.loads(str(raw_steps or "[]"))
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Question {question_id} has invalid standard tool-call JSON: {exc}"
                ) from exc
            if not isinstance(gold_steps, list) or not gold_steps:
                raise ValueError(f"Question {question_id} has no standard tool calls")
            tasks.append(
                EvalTask(
                    question_id=question_id,
                    category_id=category_id,
                    category=CATEGORY_NAMES.get(category_id, f"Category{category_id}"),
                    question=str(row[columns["题目正文"]] or ""),
                    expected_answer=str(row[columns["题目答案"]] or ""),
                    answer_type=str(row[columns["答案类型"]] or ""),
                    input_resource=str(row[columns["输入资源"]] or ""),
                    chain_mode=str(row[columns["单/多工具链调用"]] or ""),
                    gold_chain_text=str(row[columns["标准工具链调用"]] or ""),
                    gold_steps=tuple(gold_steps),
                )
            )
        return tasks
    finally:
        workbook.close()


def build_agent_question(task: EvalTask) -> str:
    if not task.input_resource:
        return task.question
    return f"{task.question}\n\nEvaluation input resource (authoritative benchmark locator; use it exactly when relevant, but retrieve and select the registered tools yourself): {task.input_resource}"
