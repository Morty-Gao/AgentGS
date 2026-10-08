"""Deterministic GeoToolEval accuracy and token-cost scoring."""

from __future__ import annotations

import json
import math
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

DEFAULT_TOKEN_BUDGET_PER_GOLD_STEP = 10000
_IGNORED_ARGUMENT_KEYS = {"_output_keys"}


def load_gold_tasks(
    workbook_path: str | Path, question_ids: Iterable[int] | None = None
) -> dict[int, dict[str, Any]]:
    """Load structured gold calls from the GeoToolEval workbook."""
    selected = None if question_ids is None else {int(value) for value in question_ids}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["评测集合"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        columns = {name: index for index, name in enumerate(headers)}
        required = {"题目序号", "所属类别", "题目正文", "题目答案", "答案类型", "标准工具调用明细"}
        missing = required - columns.keys()
        if missing:
            raise ValueError(f"GeoToolEval workbook is missing columns: {sorted(missing)}")
        tasks: dict[int, dict[str, Any]] = {}
        for row in rows:
            question_id = int(row[columns["题目序号"]])
            if selected is not None and question_id not in selected:
                continue
            raw_steps = row[columns["标准工具调用明细"]]
            try:
                gold_steps = json.loads(str(raw_steps or "[]"))
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Question {question_id} has invalid standard tool-call JSON: {exc}"
                ) from exc
            if not isinstance(gold_steps, list) or not gold_steps:
                raise ValueError(f"Question {question_id} has no standard tool calls")
            tasks[question_id] = {
                "question_id": question_id,
                "category_id": int(row[columns["所属类别"]]),
                "question": str(row[columns["题目正文"]] or ""),
                "expected_answer": str(row[columns["题目答案"]] or ""),
                "answer_type": str(row[columns["答案类型"]] or ""),
                "gold_steps": gold_steps,
            }
        if selected is not None:
            missing_ids = selected - tasks.keys()
            if missing_ids:
                raise KeyError(f"Question IDs not found in GeoToolEval: {sorted(missing_ids)}")
        return tasks
    finally:
        workbook.close()


def parse_tool_input(raw_input: Any) -> Any:
    if isinstance(raw_input, (dict, list, int, float, bool)) or raw_input is None:
        return raw_input
    text = str(raw_input).strip()
    if text.startswith("```json") and text.endswith("```"):
        text = text[7:-3].strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return text


def _clean_arguments(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            str(key): _clean_arguments(item)
            for key, item in value.items()
            if str(key) not in _IGNORED_ARGUMENT_KEYS and (not str(key).startswith("_"))
        }
    if isinstance(value, list):
        return [_clean_arguments(item) for item in value]
    return value


def _normalized_text(value: str) -> str:
    text = unicodedata.normalize("NFKC", value).strip().replace("\\", "/")
    text = re.sub("\\s+", " ", text).casefold()
    for anchor in ("geotoolagent数据集/", "benchmark/", "assets/", "outputs/geotool_runtime/"):
        index = text.find(anchor)
        if index >= 0:
            return text[index:]
    return text


def value_similarity(expected: Any, actual: Any) -> float:
    expected = _clean_arguments(expected)
    actual = _clean_arguments(actual)
    if expected is None or actual is None:
        return 1.0 if expected is actual else 0.0
    if isinstance(expected, bool) or isinstance(actual, bool):
        return 1.0 if expected is actual else 0.0
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return (
            1.0
            if math.isclose(float(expected), float(actual), rel_tol=1e-05, abs_tol=1e-08)
            else 0.0
        )
    if isinstance(expected, dict) and isinstance(actual, dict):
        keys = set(expected) | set(actual)
        if not keys:
            return 1.0
        matched = sum(
            (value_similarity(expected[key], actual[key]) for key in set(expected) & set(actual))
        )
        return 2.0 * matched / (len(expected) + len(actual))
    if isinstance(expected, list) and isinstance(actual, list):
        if not expected and (not actual):
            return 1.0
        remaining = list(actual)
        matched = 0.0
        for expected_item in expected:
            if not remaining:
                break
            similarities = [value_similarity(expected_item, item) for item in remaining]
            best_index = max(range(len(similarities)), key=similarities.__getitem__)
            matched += similarities[best_index]
            remaining.pop(best_index)
        return 2.0 * matched / (len(expected) + len(actual))
    expected_text = _normalized_text(str(expected))
    actual_text = _normalized_text(str(actual))
    if expected_text == actual_text:
        return 1.0
    if "/" in expected_text and "/" in actual_text:
        return 1.0 if expected_text.rsplit("/", 1)[-1] == actual_text.rsplit("/", 1)[-1] else 0.0
    return 0.0


def _lcs_alignment(expected: list[str], actual: list[str]) -> list[tuple[int, int]]:
    rows = len(expected) + 1
    cols = len(actual) + 1
    table = [[0] * cols for _ in range(rows)]
    for i in range(1, rows):
        for j in range(1, cols):
            if expected[i - 1] == actual[j - 1]:
                table[i][j] = table[i - 1][j - 1] + 1
            else:
                table[i][j] = max(table[i - 1][j], table[i][j - 1])
    alignment: list[tuple[int, int]] = []
    i, j = (len(expected), len(actual))
    while i > 0 and j > 0:
        if expected[i - 1] == actual[j - 1]:
            alignment.append((i - 1, j - 1))
            i -= 1
            j -= 1
        elif table[i - 1][j] >= table[i][j - 1]:
            i -= 1
        else:
            j -= 1
    return list(reversed(alignment))


def _parse_output_dict(raw_output: Any) -> dict[str, Any]:
    if isinstance(raw_output, dict):
        return raw_output
    try:
        parsed = json.loads(str(raw_output or ""))
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _effective_expected_arguments(
    gold_step: dict[str, Any], actual_arguments: Any, prior_calls: list[dict[str, Any]]
) -> tuple[Any, Any]:
    expected = _clean_arguments(gold_step.get("arguments") or {})
    actual = _clean_arguments(actual_arguments)
    if not isinstance(actual, dict):
        if isinstance(expected, dict) and len(expected) == 1:
            only_key = next(iter(expected))
            actual = {only_key: actual}
        else:
            return (expected, actual)
    prior_outputs: dict[str, Any] = {}
    prior_argument_keys: set[str] = set()
    for call in prior_calls:
        prior_outputs.update(_parse_output_dict(call.get("tool_output")))
        prior_arguments = parse_tool_input(call.get("tool_input"))
        if isinstance(prior_arguments, dict):
            prior_argument_keys.update((str(key) for key in prior_arguments))
    dynamic_keys: list[str] = []
    for consumed in gold_step.get("consumes") or []:
        key = str(consumed)
        if key in actual and key in prior_outputs:
            expected[key] = prior_outputs[key]
            dynamic_keys.append(key)
    if dynamic_keys:
        for key in list(expected):
            if key not in actual and key in prior_argument_keys:
                expected.pop(key)
    return (expected, actual)


def _tool_selection_scores(
    gold_tools: list[str], actual_tools: list[str]
) -> tuple[float, float, float, int]:
    gold_counts = Counter(gold_tools)
    actual_counts = Counter(actual_tools)
    matched = sum((min(count, actual_counts[name]) for name, count in gold_counts.items()))
    precision = matched / len(actual_tools) if actual_tools else 0.0
    recall = matched / len(gold_tools) if gold_tools else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return (precision, recall, f1, matched)


def score_task(
    *,
    gold_steps: list[dict[str, Any]],
    trace: list[dict[str, Any]],
    answer_score: float,
    token_usage: dict[str, Any] | None,
    token_budget_per_gold_step: int = DEFAULT_TOKEN_BUDGET_PER_GOLD_STEP,
) -> dict[str, Any]:
    """Score one task on the 0-100 accuracy, token-cost and final scales."""
    gold_steps = list(gold_steps or [])
    gold_tools = [str(step.get("tool_name") or "") for step in gold_steps]
    attempts = [event for event in trace or [] if event.get("type") == "tool_call"]
    successful = [event for event in attempts if event.get("status") == "success"]
    actual_tools = [str(event.get("tool_name") or "") for event in successful]
    precision, recall, tsf1, matched_success = _tool_selection_scores(gold_tools, actual_tools)
    alignment = _lcs_alignment(gold_tools, actual_tools)
    tos = len(alignment) / len(gold_tools) if gold_tools else 0.0
    parameter_scores = [0.0] * len(gold_steps)
    for gold_index, actual_index in alignment:
        actual_arguments = parse_tool_input(successful[actual_index].get("tool_input"))
        expected, actual = _effective_expected_arguments(
            gold_steps[gold_index], actual_arguments, successful[:actual_index]
        )
        parameter_scores[gold_index] = value_similarity(expected, actual)
    pas = sum(parameter_scores) / len(gold_steps) if gold_steps else 0.0
    ecr = matched_success / len(attempts) if attempts else 0.0
    fas = min(1.0, max(0.0, float(answer_score)))
    accuracy_score = 100.0 * (0.4 * fas + 0.2 * tsf1 + 0.1 * tos + 0.2 * pas + 0.1 * ecr)
    task_success = (
        math.isclose(fas, 1.0)
        and math.isclose(tsf1, 1.0)
        and math.isclose(tos, 1.0)
        and math.isclose(pas, 1.0)
        and (matched_success == len(gold_steps))
    )
    token_usage = token_usage if isinstance(token_usage, dict) else {}
    usage_complete = token_usage.get("usage_complete") is True
    total_tokens = int(token_usage.get("total_tokens") or 0)
    token_budget = max(1, int(token_budget_per_gold_step)) * max(1, len(gold_steps))
    mte: float | None
    cost_score: float | None
    final_score: float | None
    if not usage_complete:
        mte = None
        cost_score = None
        final_score = None
    else:
        mte = 0.0 if not task_success else min(1.0, token_budget / max(total_tokens, 1))
        tce_for_cost = (
            0.0 if not task_success else min(1.0, len(gold_steps) / max(len(attempts), 1))
        )
        cost_score = 100.0 * (0.6 * mte + 0.4 * tce_for_cost)
        final_score = 0.85 * accuracy_score + 0.15 * cost_score
    tce = 0.0 if not task_success else min(1.0, len(gold_steps) / max(len(attempts), 1))
    return {
        "fas": round(fas, 6),
        "tool_precision": round(precision, 6),
        "tool_recall": round(recall, 6),
        "tsf1": round(tsf1, 6),
        "lcs_length": len(alignment),
        "tos": round(tos, 6),
        "parameter_step_scores": [round(value, 6) for value in parameter_scores],
        "pas": round(pas, 6),
        "matched_successful_calls": matched_success,
        "physical_tool_call_attempts": len(attempts),
        "ecr": round(ecr, 6),
        "accuracy_score": round(accuracy_score, 4),
        "task_success_gate": task_success,
        "model_total_tokens": total_tokens,
        "token_budget": token_budget,
        "token_usage_complete": usage_complete,
        "mte": None if mte is None else round(mte, 6),
        "gold_tool_cost_units": len(gold_steps),
        "actual_tool_cost_units": len(attempts),
        "tce": round(tce, 6),
        "cost_score": None if cost_score is None else round(cost_score, 4),
        "final_score": None if final_score is None else round(final_score, 4),
        "score_ready": final_score is not None,
    }


def _domain_macro_average(samples: list[dict[str, Any]], field: str) -> float | None:
    by_domain: dict[int, list[float]] = {}
    for sample in samples:
        value = (sample.get("scores") or {}).get(field)
        if value is None:
            return None
        by_domain.setdefault(int(sample["category_id"]), []).append(float(value))
    if not by_domain:
        return None
    domain_means = [sum(values) / len(values) for values in by_domain.values()]
    return sum(domain_means) / len(domain_means)


def aggregate_scores(samples: list[dict[str, Any]]) -> dict[str, Any]:
    fields = (
        "fas",
        "tsf1",
        "tos",
        "pas",
        "ecr",
        "accuracy_score",
        "mte",
        "tce",
        "cost_score",
        "final_score",
    )
    macro = {
        field: None if (value := _domain_macro_average(samples, field)) is None else round(value, 4)
        for field in fields
    }
    return {
        "aggregation": "macro average over represented GeoToolEval domains",
        "domain_count": len({int(sample["category_id"]) for sample in samples}),
        "sample_count": len(samples),
        "token_budget_per_gold_step": (samples[0].get("scores") or {}).get("token_budget", 0)
        // max(1, len(samples[0].get("gold_steps") or []))
        if samples
        else DEFAULT_TOKEN_BUDGET_PER_GOLD_STEP,
        "total_model_tokens": sum(
            (int((sample.get("scores") or {}).get("model_total_tokens") or 0) for sample in samples)
        ),
        "score_ready": bool(samples)
        and all((bool((sample.get("scores") or {}).get("score_ready")) for sample in samples)),
        **macro,
    }
