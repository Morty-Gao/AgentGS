"""Preserve the latest released workbook and public contracts as distinct versions."""

import argparse
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook


def digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def preserve(source, destination):
    fingerprint = digest(source)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and digest(destination) != fingerprint:
        raise ValueError(f"Keep both versions; destination differs: {destination}")
    if not destination.exists():
        shutil.copy2(source, destination)
    return {
        "path": str(destination),
        "source": str(source),
        "sha256": fingerprint,
        "bytes": destination.stat().st_size,
    }


def main():
    root = next(p for p in Path(__file__).resolve().parents if (p / "pyproject.toml").exists())
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--release", type=Path)
    parser.add_argument(
        "--contracts-source",
        type=Path,
        help="Newer public-contract tree, independently of the frozen algorithm source",
    )
    args = parser.parse_args()
    source = args.source.resolve()
    release = args.release or source / "outputs/answer_format_compat_v11_20261004_side_01a10487"
    workbook = release / "GeoToolEval_固定回答格式.xlsx"
    contracts_source = args.contracts_source.resolve() if args.contracts_source else source
    latest = contracts_source / "data/answer_formats/geotooleval_v1_2.source.jsonl"
    active = contracts_source / "data/answer_formats/geotooleval_v1_1.jsonl"
    if digest(latest) != digest(active):
        raise ValueError(
            "Main-project active contracts differ from the v1.2 source; review actual versions"
        )
    contracts = [
        json.loads(line) for line in latest.read_text("utf-8").splitlines() if line.strip()
    ]
    legacy = load_workbook(
        root / "data/geotooleval/GeoToolEval.xlsx", read_only=True, data_only=True
    )
    released = load_workbook(workbook, read_only=True, data_only=True)
    try:
        old = list(legacy["评测集合"].iter_rows(values_only=True))
        new = list(released["评测集合"].iter_rows(values_only=True))
        if (
            len(old) != 604
            or len(new) != 604
            or any(tuple(a) != tuple(b[:9]) for a, b in zip(old, new, strict=True))
        ):
            raise ValueError("Released tasks differ from the frozen original nine columns")
        by_id = {int(row[0]): row for row in new[1:]}
        if len(contracts) != 603 or len(by_id) != 603:
            raise ValueError("Contract/workbook coverage must be 603 unique tasks")
        for row in contracts:
            if row["question"] != by_id[row["question_id"]][2]:
                raise ValueError(f"Public question mismatch: {row['question_id']}")
    finally:
        legacy.close()
        released.close()
    records = []
    contract_path = "data/answer_formats/geotooleval_v1_2." + digest(latest)[:12] + ".jsonl"
    system_path = "data/answer_formats/geotooleval_v1_2." + digest(latest)[:12] + "_system.txt"
    records.append(preserve(workbook, root / "data/geotooleval/GeoToolEval_固定回答格式.xlsx"))
    records.append(preserve(latest, root / contract_path))
    records.append(
        preserve(
            contracts_source / "data/answer_formats/geotooleval_v1_1_system.txt",
            root / system_path,
        )
    )
    records.append(
        preserve(release / "answer_formats.jsonl", root / "data/answer_formats/archive/v1_1.jsonl")
    )
    records.append(
        preserve(
            release / "answer_system_prompt.txt",
            root / "data/answer_formats/archive/v1_1_system.txt",
        )
    )
    for name in (
        "delivery_manifest.json",
        "delivery_validation.json",
        "upgrade_manifest.json",
        "rescore_and_replay.json",
    ):
        records.append(preserve(release / name, root / "data/geotooleval/release_v1_1" / name))
    manifest = {
        "schema_version": 1,
        "selected_workbook": "data/geotooleval/GeoToolEval_固定回答格式.xlsx",
        "public_contracts": contract_path,
        "public_system_prompt": system_path,
        "contract_revisions": sorted(
            {row["answer_format"].get("revision", "unknown") for row in contracts}
        ),
        "compatibility_contracts": "data/answer_formats/archive/v1_1.jsonl",
        "question_count": 603,
        "original_nine_columns_identical": True,
        "workbook_format_version": "1.1",
        "active_contract_version": "1.2",
        "version_note": "Released workbook retains v1.1 columns; newer main-project v1.2 public contracts govern runtime prompts. Original FAS remains a separate frozen protocol.",
        "records": [{**r, "path": str(Path(r["path"]).relative_to(root))} for r in records],
    }
    path = root / "data/geotooleval/version_manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"manifest": str(path), "tasks": 603, "active_contract_version": "1.2"}))


if __name__ == "__main__":
    main()
